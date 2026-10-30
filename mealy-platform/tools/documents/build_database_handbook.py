from __future__ import annotations

import argparse
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ORANGE = "FF6B18"
ORANGE_DARK = "D94D00"
INK = "171717"
SLATE = "4A4A4A"
MUTED = "6B6B6B"
PALE_ORANGE = "FFF2E8"
PALE_GRAY = "F6F6F4"
MID_GRAY = "D9D9D9"
WHITE = "FFFFFF"

SKIP_COLUMNS = {"id", "created_at", "updated_at"}

DOMAIN_BY_MIGRATION = {
    "001_foundation.sql": "Identity tenant and platform control",
    "002_catalog_pricing_inventory.sql": "Catalog pricing tax and inventory",
    "003_orders_tokens_kot.sql": "Orders tokens and kitchen operations",
    "004_payments_promotions_wallet.sql": "Payments promotions wallet and finance",
    "005_pos_admin_operations.sql": "POS administration and service operations",
    "006_reliability_audit.sql": "Reliability security and audit",
}

PURPOSE_OVERRIDES = {
    "schema_metadata": "Records schema identity and database-level metadata.",
    "app_users": "Stores the global user identity shared by students staff and administrators.",
    "auth_identities": "Links users to login providers and normalized external identifiers.",
    "user_sessions": "Stores revocable authenticated sessions and their security lifecycle.",
    "campuses": "Defines campuses managed by the central platform.",
    "canteens": "Defines canteen operating organizations independent of any one campus.",
    "campus_canteens": "Assigns a canteen organization to a campus and anchors physical outlets.",
    "outlets": "Represents a physical canteen outlet at a campus assignment.",
    "service_counters": "Defines pickup or service counters within an outlet.",
    "pos_terminals": "Registers POS devices and binds them to outlet and counter scope.",
    "student_profiles": "Adds student-specific attributes to the global user account.",
    "student_campus_memberships": "Tracks student association or verification for a campus.",
    "access_memberships": "Grants a user membership at platform campus canteen or outlet scope.",
    "membership_roles": "Assigns scoped roles to access memberships.",
    "canteen_features": "Controls optional capabilities independently for each canteen.",
    "menus": "Defines canteen menus and their publication lifecycle.",
    "menu_outlets": "Publishes a menu to one or more permitted outlets.",
    "catalog_categories": "Stores hierarchical food and beverage categories.",
    "catalog_items": "Stores canonical sellable items and their food metadata.",
    "item_variants": "Stores sizes or other sellable variants including the default variant.",
    "modifier_groups": "Defines add-on groups with selection and quantity rules.",
    "modifier_options": "Defines individual add-ons and their incremental price.",
    "item_modifier_groups": "Assigns valid add-on groups to an item or a specific variant.",
    "outlet_item_settings": "Overrides item price availability and sold-out state per outlet.",
    "outlet_variant_settings": "Overrides variant price and availability per outlet.",
    "price_rules": "Stores scheduled channel-aware and outlet-aware price adjustments.",
    "tax_categories": "Groups catalog items into tax treatment categories.",
    "tax_rates": "Stores effective-dated tax components including inclusive and compound rules.",
    "tax_category_rates": "Maps tax categories to their component rates.",
    "service_charges": "Stores fixed or percentage service charges by scope and channel.",
    "inventory_items": "Defines ingredients or stock-managed supplies.",
    "recipes": "Maps catalog variants to inventory consumption quantities.",
    "inventory_balances": "Stores the current outlet inventory balance with versioning.",
    "inventory_movements": "Provides the immutable inventory movement ledger.",
    "carts": "Stores an active or converted cart for student or POS checkout.",
    "cart_items": "Stores cart item selections before an order snapshot is created.",
    "cart_item_modifiers": "Stores selected cart add-ons and quantities.",
    "order_number_sequences": "Allocates atomic daily order numbers per outlet.",
    "token_sequences": "Allocates atomic daily pickup tokens per service counter.",
    "orders": "Stores the authoritative order identity totals payment state and fulfillment state.",
    "order_items": "Stores immutable item and variant snapshots accepted at checkout.",
    "order_item_modifiers": "Stores immutable selected add-on snapshots for each order item.",
    "order_tax_lines": "Stores tax component snapshots and calculated amounts for an order.",
    "order_charge_lines": "Stores fee and service-charge snapshots for an order.",
    "order_status_events": "Records the append-only fulfillment status history.",
    "order_tokens": "Links an order to its counter-scoped pickup token.",
    "kitchen_stations": "Defines kitchen preparation destinations within an outlet.",
    "item_station_routes": "Routes catalog items to appropriate kitchen stations.",
    "kots": "Stores Kitchen Order Tickets and revision lifecycle.",
    "kot_items": "Stores KOT item snapshots and preparation status.",
    "kot_item_modifiers": "Stores KOT add-on snapshots for kitchen visibility.",
    "kot_status_events": "Records KOT lifecycle changes.",
    "payment_method_configs": "Configures enabled payment methods per canteen or outlet using secret references.",
    "payment_intents": "Represents an amount attempted through one payment method for an order.",
    "payment_transactions": "Stores immutable authorization capture sale refund and adjustment events.",
    "refunds": "Tracks order refunds with idempotency and provider references.",
    "wallets": "Stores a student's versioned balance by currency.",
    "wallet_transactions": "Provides the immutable idempotent wallet ledger.",
    "promotions": "Defines eligibility and value rules for canteen promotions.",
    "promotion_codes": "Defines redeemable codes associated with promotions.",
    "order_discount_lines": "Stores immutable discount snapshots accepted on an order.",
    "promotion_redemptions": "Tracks promotion consumption and reversal per order.",
    "platform_collections": "Calculates platform and canteen settlement amounts for an order.",
    "payouts": "Tracks payouts from the platform to a canteen.",
    "payout_collections": "Allocates collected order proceeds into a payout.",
    "printers": "Registers outlet printers and connection configuration.",
    "printer_routes": "Routes document types or stations to printers.",
    "print_jobs": "Queues retryable KOT and receipt print work.",
    "cash_registers": "Defines cash drawers or registers at an outlet.",
    "pos_shifts": "Tracks register opening closing and reconciliation.",
    "cash_movements": "Records cash paid in paid out and corrective movements.",
    "receipts": "Stores issued receipt references and order linkage.",
    "feedback": "Stores an order-level student feedback submission.",
    "feedback_items": "Stores item-specific ratings or comments within feedback.",
    "push_devices": "Registers notification-capable student devices.",
    "notification_preferences": "Stores student notification-channel choices.",
    "notifications": "Tracks notification content delivery state and retries.",
    "api_idempotency_records": "Prevents duplicate processing of retryable API operations.",
    "webhook_events": "Protects provider webhook processing against replay and duplicate delivery.",
    "outbox_events": "Queues events transactionally for reliable downstream publication.",
    "audit_logs": "Records security-sensitive and administrative activity.",
    "data_subject_requests": "Tracks privacy access correction or deletion requests.",
    "login_challenges": "Stores expiring OTP or login verification challenges.",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_run_font(run, name="Arial", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction: str, display: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def crop_logo(source: Path, output: Path) -> None:
    image = Image.open(source).convert("RGBA")
    rgb = image.convert("RGB")
    mask = Image.new("L", rgb.size, 0)
    pixels = rgb.load()
    mask_pixels = mask.load()
    for y in range(rgb.height):
        for x in range(rgb.width):
            r, g, b = pixels[x, y]
            if r > 180 and g < 170 and b < 120:
                mask_pixels[x, y] = 255
    bbox = mask.getbbox()
    if bbox is None:
        raise ValueError("Logo does not contain detectable orange artwork")
    pad = 24
    left = max(0, bbox[0] - pad)
    top = max(0, bbox[1] - pad)
    right = min(image.width, bbox[2] + pad)
    bottom = min(image.height, bbox[3] + pad)
    cropped = image.crop((left, top, right, bottom))
    output.parent.mkdir(parents=True, exist_ok=True)
    cropped.save(output)


def load_font(size: int, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def rounded_box(draw, xy, fill, outline=None, radius=28, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def draw_arrow(draw, start, end, color=(90, 90, 90), width=5):
    draw.line([start, end], fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 18 * direction, ey - 11), (ex - 18 * direction, ey + 11)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 11, ey - 18 * direction), (ex + 11, ey - 18 * direction)]
    draw.polygon(points, fill=color)


def centered_multiline(draw, box, text, font, fill=(23, 23, 23), spacing=6):
    left, top, right, bottom = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(((left + right - width) / 2, (top + bottom - height) / 2), text, font=font, fill=fill, spacing=spacing, align="center")


def create_hierarchy_diagram(output: Path) -> None:
    canvas = Image.new("RGB", (1500, 790), "white")
    draw = ImageDraw.Draw(canvas)
    title = load_font(42, True)
    label = load_font(29, True)
    small = load_font(24)
    draw.text((70, 45), "Tenant and location hierarchy", font=title, fill=(23, 23, 23))
    boxes = [
        (60, 220, 300, 375, "Central\nplatform"),
        (350, 220, 590, 375, "Campus"),
        (640, 220, 880, 375, "Canteen\norganization"),
        (930, 220, 1170, 375, "Campus canteen\nassignment"),
        (1220, 220, 1460, 375, "Physical\noutlet"),
    ]
    for index, (l, t, r, b, text) in enumerate(boxes):
        fill = (255, 107, 24) if index in (0, 2) else (255, 242, 232)
        outline = (217, 77, 0) if index in (0, 2) else (255, 154, 97)
        rounded_box(draw, (l, t, r, b), fill, outline)
        centered_multiline(draw, (l, t, r, b), text, label, fill=(255, 255, 255) if index in (0, 2) else (23, 23, 23))
        if index < len(boxes) - 1:
            draw_arrow(draw, (r + 8, 298), (boxes[index + 1][0] - 8, 298))
    lower = [
        (965, 510, 1135, 635, "Service\ncounter"),
        (1245, 510, 1415, 635, "POS\nterminal"),
    ]
    draw_arrow(draw, (1340, 390), (1060, 500))
    draw_arrow(draw, (1145, 573), (1235, 573))
    for l, t, r, b, text in lower:
        rounded_box(draw, (l, t, r, b), (246, 246, 244), (190, 190, 190), radius=22)
        centered_multiline(draw, (l, t, r, b), text, small)
    draw.text((70, 705), "A canteen may operate on multiple campuses while each outlet retains explicit campus and counter scope.", font=small, fill=(74, 74, 74))
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, quality=95)


def create_order_diagram(output: Path) -> None:
    canvas = Image.new("RGB", (1500, 850), "white")
    draw = ImageDraw.Draw(canvas)
    title = load_font(42, True)
    label = load_font(28, True)
    small = load_font(23)
    draw.text((70, 45), "Authoritative pickup order flow", font=title, fill=(23, 23, 23))
    stages = ["Browse\nmenu", "Server\nquote", "Create\norder", "Assign\ntoken", "Route\nKOT", "Prepare", "Ready", "Collect"]
    x = 55
    boxes = []
    for i, stage in enumerate(stages):
        w = 150 if i not in (1, 2, 4) else 165
        box = (x, 215, x + w, 345)
        boxes.append(box)
        fill = (255, 107, 24) if i in (2, 6) else (255, 242, 232)
        rounded_box(draw, box, fill, (217, 77, 0), radius=24, width=3)
        centered_multiline(draw, box, stage, label, fill=(255, 255, 255) if i in (2, 6) else (23, 23, 23))
        x += w + 28
    for a, b in zip(boxes, boxes[1:]):
        draw_arrow(draw, (a[2] + 3, 280), (b[0] - 5, 280), width=4)
    status = ["placed", "confirmed", "preparing", "ready for pickup", "collected"]
    x_positions = [275, 500, 725, 950, 1190]
    draw.text((70, 475), "Seeded status lifecycle", font=label, fill=(23, 23, 23))
    for i, (text, cx) in enumerate(zip(status, x_positions)):
        draw.ellipse((cx - 28, 565, cx + 28, 621), fill=(255, 107, 24), outline=(217, 77, 0), width=3)
        tw = draw.textbbox((0, 0), text, font=small)[2]
        draw.text((cx - tw / 2, 650), text, font=small, fill=(74, 74, 74))
        if i < len(status) - 1:
            draw_arrow(draw, (cx + 35, 593), (x_positions[i + 1] - 35, 593), width=4)
    draw.text((70, 770), "Order totals status token KOT and events remain server-owned and are shared by student POS and admin views.", font=small, fill=(74, 74, 74))
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, quality=95)


def create_split_diagram(output: Path) -> None:
    canvas = Image.new("RGB", (1500, 900), "white")
    draw = ImageDraw.Draw(canvas)
    title = load_font(42, True)
    label = load_font(29, True)
    small = load_font(23)
    draw.text((70, 45), "Split payment data model", font=title, fill=(23, 23, 23))
    order = (585, 150, 915, 290)
    rounded_box(draw, order, (255, 107, 24), (217, 77, 0), radius=30)
    centered_multiline(draw, order, "One order\npartially paid or paid", label, fill=(255, 255, 255))
    tenders = [
        (95, 445, 385, 585, "Payment intent A\nCash amount"),
        (605, 445, 895, 585, "Payment intent B\nUPI or card amount"),
        (1115, 445, 1405, 585, "Payment intent C\nWallet amount"),
    ]
    for box in tenders:
        draw_arrow(draw, (750, 305), ((box[0] + box[2]) // 2, box[1] - 10), width=5)
        rounded_box(draw, box[:4], (255, 242, 232), (255, 154, 97), radius=28)
        centered_multiline(draw, box[:4], box[4], label)
    txn_y = 715
    for box in tenders:
        cx = (box[0] + box[2]) // 2
        draw_arrow(draw, (cx, box[3] + 8), (cx, txn_y - 10), width=4)
        rounded_box(draw, (cx - 120, txn_y, cx + 120, txn_y + 95), (246, 246, 244), (190, 190, 190), radius=18)
        centered_multiline(draw, (cx - 120, txn_y, cx + 120, txn_y + 95), "Immutable\ntransaction", small)
    draw.text((70, 850), "Storage supports multiple tenders. Atomic remaining-balance and overpayment enforcement belongs in the backend and is still TBD.", font=small, fill=(74, 74, 74))
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, quality=95)


def parse_tables(migrations_dir: Path):
    tables = []
    pattern = re.compile(r"CREATE TABLE\s+(\w+)\s*\((.*?)\n\);", re.S | re.I)
    for sql_path in sorted(migrations_dir.glob("*.sql")):
        if sql_path.name.startswith("007_"):
            continue
        sql = sql_path.read_text(encoding="utf-8")
        for match in pattern.finditer(sql):
            name, body = match.group(1), match.group(2)
            columns = []
            for line in body.splitlines():
                stripped = line.strip().rstrip(",")
                if not stripped or stripped.upper().startswith(("PRIMARY KEY", "UNIQUE", "CHECK", "CONSTRAINT", "FOREIGN KEY")):
                    continue
                column = stripped.split()[0].strip('"')
                if re.match(r"^[a-z][a-z0-9_]*$", column):
                    columns.append(column)
            references = []
            for ref in re.findall(r"REFERENCES\s+(\w+)\s*\(", body, re.I):
                if ref not in references:
                    references.append(ref)
            tables.append({
                "name": name,
                "migration": sql_path.name,
                "domain": DOMAIN_BY_MIGRATION.get(sql_path.name, "Other"),
                "columns": columns,
                "references": references,
            })
    return tables


def humanize(name: str) -> str:
    return name.replace("_", " ").strip().capitalize()


def purpose_for(name: str) -> str:
    if name in PURPOSE_OVERRIDES:
        return PURPOSE_OVERRIDES[name]
    if name.endswith("_settings"):
        return f"Stores scoped configuration for {humanize(name[:-9]).lower()}."
    if name.endswith("_events"):
        return f"Records append-only lifecycle events for {humanize(name[:-7]).lower()}."
    if name.endswith("_sequences"):
        return f"Provides concurrency-safe numbering for {humanize(name[:-10]).lower()}."
    if name.endswith("_items"):
        return f"Stores item-level records belonging to {humanize(name[:-6]).lower()}."
    if name.endswith("_roles"):
        return f"Maps role assignments for {humanize(name[:-6]).lower()}."
    return f"Stores {humanize(name).lower()} records used by the Mealy platform."


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.14

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(31)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 21, 18, 9),
        ("Heading 2", 15, 14, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)


def add_header_footer(section, logo_path: Path) -> None:
    header = section.header
    header.is_linked_to_previous = False
    table = header.add_table(rows=1, cols=2, width=Inches(6.85))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(5.75)
    table.columns[1].width = Inches(1.1)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = left.add_run("MEALY DATABASE HANDBOOK")
    set_run_font(run, size=8, bold=True, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    picture = right.add_run().add_picture(str(logo_path), width=Inches(0.72))
    picture._inline.docPr.set("descr", "Mealy logo")
    for cell in table.rows[0].cells:
        set_cell_margins(cell, 0, 0, 0, 0)
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Mealy platform database  •  ")
    set_run_font(run, size=8, color=MUTED)
    add_field(p, "PAGE", "1")
    for run in p.runs:
        set_run_font(run, size=8, color=MUTED)


def add_cover(doc: Document, logo_path: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(84)
    picture = p.add_run().add_picture(str(logo_path), width=Inches(3.55))
    picture._inline.docPr.set("descr", "Mealy logo")

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.add_run("Mealy Database Handbook")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Production foundation for student ordering canteen POS administration and shared services")
    set_run_font(run, size=13, color=SLATE)
    p.paragraph_format.space_after = Pt(42)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = [("PostgreSQL", "16 plus"), ("Schema", "112 tables"), ("Status", "Verified foundation")]
    widths = [1.85, 1.85, 2.2]
    for i, ((label, value), width) in enumerate(zip(labels, widths)):
        cell = table.cell(0, i)
        cell.width = Inches(width)
        set_cell_shading(cell, PALE_ORANGE if i != 1 else PALE_GRAY)
        set_cell_margins(cell, 180, 160, 180, 160)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label.upper() + "\n")
        set_run_font(r, size=8, bold=True, color=MUTED)
        r = p.add_run(value)
        set_run_font(r, size=13, bold=True, color=INK)
    set_table_borders(table, color="FFFFFF", size="0")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(78)
    r = p.add_run("Prepared 14 September 2026")
    set_run_font(r, size=9.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Architecture reference and implementation guide")
    set_run_font(r, size=9.5, color=MUTED)
    doc.add_page_break()


def add_intro(doc: Document) -> None:
    doc.add_heading("Purpose and main conclusion", level=1)
    doc.add_paragraph(
        "This handbook explains the complete PostgreSQL database foundation created for Mealy. It is intended for product owners, backend engineers, POS integrators, administrators, testers, and operators who need one readable reference for what the schema covers and what remains to be implemented in application code."
    )
    p = doc.add_paragraph()
    r = p.add_run("Main conclusion. ")
    set_run_font(r, bold=True)
    p.add_run(
        "The database is a verified production-oriented foundation for a multi-campus food ordering platform. It covers student ordering, canteen POS, canteen and central administration, catalog customization, taxes, KOTs, tokens, payments, wallet, printing, audit, and reliability. A passing schema does not by itself make the product production-ready; backend authorization, calculations, provider integrations, concurrency tests, and policy decisions still need implementation or approval."
    )

    doc.add_heading("Document map", level=2)
    items = [
        "Database at a glance",
        "Split payment capability",
        "System boundaries and tenant model",
        "Catalog pricing tax and inventory",
        "Orders tokens and kitchen operations",
        "Payments promotions wallet and settlement",
        "POS administration reliability and security",
        "Data invariants migrations and verification",
        "Production readiness and TBD decisions",
        "Complete table function and trigger catalogue",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_metrics(doc: Document) -> None:
    doc.add_heading("Database at a glance", level=1)
    doc.add_paragraph(
        "The foundation uses framework-neutral append-only SQL migrations for PostgreSQL 16 and later. Monetary values use integer minor units, identifiers use UUIDs, timestamps are timezone-aware, and business events preserve immutable snapshots so historical orders do not change when a menu or tax rule changes."
    )
    rows = [
        ("Tables", "112", "Domain and operational records"),
        ("Foreign keys", "215", "Relationship and tenant integrity"),
        ("Check constraints", "305", "Allowed states ranges and value shapes"),
        ("Indexes", "220", "Lookup uniqueness and operational performance"),
        ("Functions", "11", "Scope checks numbering tokens transitions and event recording"),
        ("Non internal triggers", "74", "Validation events and automatic updated timestamps"),
        ("Migrations", "7", "Ordered schema seed function and constraint layers"),
    ]
    add_data_table(doc, ["Measure", "Verified count", "Purpose"], rows, [1.55, 1.25, 3.75])

    doc.add_heading("Four product surfaces", level=2)
    rows = [
        ("Student ordering", "Identity campus context menu cart checkout order tracking token collection wallet and feedback models"),
        ("Canteen POS", "Counter orders incoming orders KOT status payment receipt printer register shift and cash models"),
        ("Canteen and central admin", "Campuses canteens outlets staff roles features settings onboarding subscriptions and audit models"),
        ("Shared backend", "Pricing tax order identity sequences idempotency webhooks outbox and authoritative state transitions"),
    ]
    add_data_table(doc, ["Surface", "Database coverage"], rows, [1.75, 4.8])

    doc.add_heading("Default feature posture", level=2)
    rows = [
        ("Enabled core", "Student ordering POS KOT tokens printing cash payment and multi POS"),
        ("Modeled but disabled", "Wallet coupons feedback inventory dine in reservations delivery subscriptions commissions and online payments"),
        ("Scope rule", "A modeled table is not launch approval; optional capabilities require an explicit product decision and canteen feature enablement"),
    ]
    add_data_table(doc, ["Class", "Capabilities"], rows, [1.75, 4.8])


def add_data_table(doc: Document, headers, rows, widths=None, font_size=9.1) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, SLATE)
        set_cell_margins(cell, 120, 130, 120, 130)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(text))
        set_run_font(run, size=9, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for i, value in enumerate(values):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            if row_index % 2:
                set_cell_shading(cell, PALE_GRAY)
            set_cell_margins(cell, 105, 130, 105, 130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 1 and len(headers) == 3 and headers[1].lower().endswith("count") else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_picture(doc: Document, path: Path, width: float, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", alt)


def add_split_payment(doc: Document, diagram: Path) -> None:
    doc.add_heading("Split payment capability", level=1)
    p = doc.add_paragraph()
    r = p.add_run("Answer. ")
    set_run_font(r, bold=True)
    p.add_run(
        "Yes at the data-model level, with an important implementation qualification. An order may own multiple payment intents, each intent may use a different enabled payment method and amount, and the order payment state explicitly includes partially paid. This can represent cash plus UPI, wallet plus card, or other mixed tenders."
    )
    add_picture(doc, diagram, 6.55, "Diagram showing one order connected to multiple payment intents and transactions")

    doc.add_heading("What is already supported", level=2)
    supported = [
        "Multiple payment_intents rows may reference the same order; there is no one-intent-per-order constraint.",
        "Each intent stores its own method configuration amount currency status provider reference and idempotency key.",
        "Immutable payment_transactions preserve authorization capture sale void refund chargeback and adjustment events.",
        "orders.payment_status supports unpaid pending partially_paid paid partially_refunded refunded failed and voided.",
        "Refunds may target a specific successful payment transaction, which enables tender-aware refund routing.",
        "Wallet ledger entries can reference both an order and a payment transaction.",
    ]
    for item in supported:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("What the backend must still enforce", level=2)
    rows = [
        ("Atomic balance", "Lock or version the order while adding or capturing a tender and calculate the remaining amount server-side.", "TBD implementation"),
        ("Overpayment", "Reject or explicitly authorize payments whose successful total exceeds the payable order total.", "TBD policy"),
        ("State derivation", "Derive partially paid or paid from successful net captures rather than trusting a client-supplied status.", "TBD implementation"),
        ("Failure recovery", "Define how expired failed or reversed tenders affect the remaining amount and retry experience.", "TBD policy"),
        ("Refund allocation", "Define proportional manual or tender-priority refund allocation and ensure totals cannot exceed net captured value.", "TBD policy"),
        ("KOT timing", "Decide whether KOT creation occurs before payment after first tender or only after full settlement.", "Blocking product decision"),
    ]
    add_data_table(doc, ["Control", "Required behavior", "Status"], rows, [1.2, 4.0, 1.35], font_size=8.6)
    p = doc.add_paragraph()
    r = p.add_run("Readiness classification. ")
    set_run_font(r, bold=True)
    p.add_run("Split payment is structurally supported but not yet production-ready until these backend rules and tests are implemented.")


def add_tenant_section(doc: Document, diagram: Path) -> None:
    doc.add_heading("System boundaries and tenant model", level=1)
    add_picture(doc, diagram, 6.55, "Diagram of central platform campus canteen assignment outlet counter and POS terminal hierarchy")
    doc.add_paragraph(
        "The canteen is the operating organization. A campus_canteens record assigns that canteen to a campus. Physical outlets belong to the assignment, counters belong to an outlet, and terminals are registered at the correct outlet and counter. This avoids treating a canteen brand and a physical serving location as the same entity."
    )
    rows = [
        ("Platform", "Global settings feature definitions system roles campus onboarding and cross-tenant governance"),
        ("Campus", "Campus profile timezone contact settings and canteen assignments"),
        ("Canteen", "Operating organization catalog ownership feature entitlements payment configuration and staff scope"),
        ("Outlet", "Physical operation hours closures menu publication pricing inventory kitchen printers registers and reporting scope"),
        ("Counter", "Pickup location daily token namespace and terminal association"),
        ("Terminal", "Registered POS device operational context and audit attribution"),
    ]
    add_data_table(doc, ["Level", "Responsibility"], rows, [1.3, 5.25])

    doc.add_heading("Identity authorization and feature control", level=2)
    for item in [
        "Global app users may have provider identities revocable sessions and expiring login challenges.",
        "Student profiles and campus memberships remain distinct from staff access.",
        "Staff authorization is modeled through scoped memberships roles permissions and membership-role assignments.",
        "Feature definitions and canteen_features allow optional modules to remain disabled until approved.",
        "Final runtime database grants and row-level security depend on the selected backend authentication framework and remain TBD.",
    ]:
        doc.add_paragraph(item, style="List Bullet")


def add_catalog_section(doc: Document) -> None:
    doc.add_heading("Catalog pricing tax and inventory", level=1)
    doc.add_paragraph(
        "Catalog data separates canonical products from outlet-specific selling conditions. A menu can be scheduled and assigned to outlets, categories can be hierarchical, and each item can expose variants plus add-on groups with explicit selection rules. Historical order lines copy the accepted names quantities and prices so later catalog edits do not rewrite a receipt or KOT."
    )
    rows = [
        ("Menus and schedules", "Menus menu outlets schedules categories and category ordering", "Core"),
        ("Items and variants", "Canonical items media SKU dietary class preparation time variants and one default variant", "Core"),
        ("Add-ons", "Modifier groups options min max free selections quantities item-level and variant-level assignment", "Core"),
        ("Outlet selling", "Price overrides availability sold-out windows daily limits and channel-aware scheduled price rules", "Core"),
        ("Tax", "Tax categories multiple effective rates inclusive pricing compound order and immutable order tax lines", "Core"),
        ("Service charges", "Fixed or percentage charges by scope channel and effective period", "Supported"),
        ("Allergens", "Allergen registry and item mappings", "Supported"),
        ("Inventory", "Ingredients recipes balances versions and immutable movements", "Optional disabled"),
    ]
    add_data_table(doc, ["Capability", "Model coverage", "Default scope"], rows, [1.35, 4.15, 1.05], font_size=8.7)

    doc.add_heading("Pricing rules that application code must own", level=2)
    for item in [
        "Select the active outlet menu schedule availability and price rule for the requested channel and time.",
        "Validate variant and modifier eligibility and min/max quantities before calculating totals.",
        "Apply tax component ordering inclusive-tax extraction discounts and service charges using approved rounding fixtures.",
        "Persist only server-calculated accepted snapshots on the order.",
        "Accounting approval for exact tax promotion and rounding precedence remains TBD before API implementation.",
    ]:
        doc.add_paragraph(item, style="List Bullet")


def add_order_section(doc: Document, diagram: Path) -> None:
    doc.add_heading("Orders tokens and kitchen operations", level=1)
    add_picture(doc, diagram, 6.55, "Diagram of the authoritative Mealy pickup order flow and status lifecycle")
    doc.add_paragraph(
        "Student orders, POS-created counter orders, kiosk orders, administrative orders, and API orders use the same authoritative orders table. Each order receives an outlet-scoped daily order number. Pickup tokens are allocated atomically by service counter and business date, while KOTs route immutable item and modifier snapshots to kitchen stations."
    )
    rows = [
        ("Order identity", "UUID plus atomic daily outlet order number and optimistic version"),
        ("Financial snapshot", "Subtotal discount tax charge total currency item variant modifier and pricing snapshots"),
        ("Fulfillment", "Pickup by default with modeled dine-in and delivery options disabled unless enabled"),
        ("Token", "Counter and service-date scoped sequence with prefix and idempotent allocation"),
        ("KOT", "Multiple tickets revisions kitchen station routing item modifiers and status events"),
        ("Status", "Data-driven transitions actor restrictions reason requirements full history and outbox publication"),
    ]
    add_data_table(doc, ["Concern", "Database behavior"], rows, [1.35, 5.2])
    p = doc.add_paragraph()
    r = p.add_run("Implemented pickup baseline. ")
    set_run_font(r, bold=True)
    p.add_run("placed to confirmed to preparing to ready for pickup to collected, with controlled rejection and cancellation paths. Final UI and business approval remains TBD.")


def add_payment_section(doc: Document) -> None:
    doc.add_heading("Payments promotions wallet and platform finance", level=1)
    doc.add_paragraph(
        "Payment configuration stores only secret-manager references and safe public configuration. Provider intent identifiers and transaction identifiers are unique when present, while idempotency keys protect retryable order-scoped actions. Refunds, wallet ledger entries, commission calculations, collections, payouts, and subscriptions are separate records so reconciliation does not depend on mutable order fields alone."
    )
    rows = [
        ("Payment methods", "Cash UPI card Razorpay wallet bank transfer and other, scoped to canteen or outlet and channel"),
        ("Intents", "Order amount currency method provider lifecycle expiration failure details and idempotency"),
        ("Transactions", "Authorization capture sale void refund chargeback and adjustment events"),
        ("Refunds", "Order and optional transaction target reason provider reference status and request attribution"),
        ("Promotions", "Percentage fixed item fixed buy X get Y and free item definitions with codes and limits"),
        ("Wallet", "One balance per student and currency with versioning and immutable idempotent movements"),
        ("Platform finance", "Commission rules collections net payable payouts and payout allocation"),
        ("Subscriptions", "Plans canteen subscriptions billing lifecycle and included feature metadata"),
    ]
    add_data_table(doc, ["Area", "Coverage"], rows, [1.45, 5.1])
    doc.add_paragraph(
        "Wallet coupons commissions subscriptions and online payments are modeled but default disabled. Launch methods provider selection settlement timing and KOT timing remain product decisions."
    )


def add_pos_reliability_section(doc: Document) -> None:
    doc.add_heading("POS administration reliability and security", level=1)
    doc.add_heading("POS and canteen operations", level=2)
    rows = [
        ("Printing", "Printers routes retryable jobs error details attempts and reprint-safe document linkage"),
        ("Cash", "Registers one active shift opening and closing balances movements and reconciliation"),
        ("Receipts", "Outlet sequences receipt numbers issue status and order association"),
        ("Optional service", "Dining areas tables reservations student addresses and delivery assignments"),
        ("Engagement", "Feedback item ratings push devices notification preferences and delivery state"),
        ("Onboarding", "Canteen setup checklist ownership completion and operational readiness steps"),
    ]
    add_data_table(doc, ["Area", "Model coverage"], rows, [1.45, 5.1])

    doc.add_heading("Reliability and traceability", level=2)
    for item in [
        "API idempotency records protect retryable creates and side effects.",
        "Webhook events store provider delivery identity processing status attempts and redacted payloads.",
        "Transactional outbox events are committed with domain changes and published later with retries.",
        "Audit logs retain actor scope request metadata action target and structured before/after context.",
        "Data-subject requests model privacy access correction export and deletion workflows.",
        "Login challenges separate expiring verification state from long-lived sessions.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Security boundary", level=2)
    doc.add_paragraph(
        "Browsers and POS clients must never receive direct authority to calculate totals change payment state or mutate arbitrary tenant rows. The backend must derive scope from the authenticated principal, recalculate prices, verify payment signatures and webhooks, invoke database functions through least-privilege roles, and publish updates only after the transaction commits."
    )


def add_invariants(doc: Document) -> None:
    doc.add_heading("Data standards and enforced invariants", level=1)
    rows = [
        ("Identifiers", "UUID primary keys for distributed-safe entity identity; bigserial only for append-only event ordering where appropriate"),
        ("Money", "Integer minor units plus ISO three-character currency code; no floating-point money"),
        ("Inventory", "numeric 18 by 6 for precise ingredient quantities"),
        ("Time", "Timezone-aware timestamps and explicit business date for order token register and receipt sequences"),
        ("History", "Immutable order KOT tax charge discount payment inventory wallet and audit records"),
        ("Concurrency", "Atomic sequence functions idempotency keys unique provider IDs optimistic versions and row locks where required"),
        ("Tenancy", "Foreign keys scope-validation triggers and explicit campus canteen outlet counter relationships"),
        ("Secrets", "Only secret references are persisted in payment configuration; payloads are redacted before storage"),
        ("Lifecycle", "Check constraints plus data-driven status transitions and actor restrictions"),
    ]
    add_data_table(doc, ["Standard", "Implementation"], rows, [1.35, 5.2])

    doc.add_heading("Database functions", level=2)
    rows = [
        ("set_updated_at", "Maintains updated timestamps automatically."),
        ("validate_terminal_counter_scope", "Prevents terminal and counter scope mismatch."),
        ("validate_order_scope", "Validates campus canteen outlet counter terminal and actor relationships on orders."),
        ("validate_membership_role_scope", "Prevents assigning a role outside a membership's scope."),
        ("validate_menu_outlet_scope", "Prevents publishing a canteen menu to an unrelated outlet."),
        ("validate_item_modifier_scope", "Prevents cross-canteen item and modifier assignment."),
        ("validate_promotion_code_scope", "Prevents cross-canteen promotion and code relationships."),
        ("next_order_number", "Allocates a concurrency-safe daily outlet order number."),
        ("allocate_order_token", "Allocates an idempotent counter and date scoped token."),
        ("transition_order", "Applies authorized version-aware order status transitions and history."),
        ("record_new_order", "Creates the initial status event and transactional outbox record."),
    ]
    add_data_table(doc, ["Function", "Responsibility"], rows, [2.15, 4.4], font_size=8.7)


def add_migrations_testing(doc: Document) -> None:
    doc.add_heading("Migrations operations and verification", level=1)
    doc.add_heading("Migration layers", level=2)
    rows = [
        ("001", "Foundation", "Identity tenancy RBAC settings hours closures and feature control"),
        ("002", "Catalog", "Menus categories items variants modifiers pricing taxes allergens and inventory"),
        ("003", "Orders", "Carts orders snapshots sequences tokens stations KOTs and events"),
        ("004", "Payments", "Methods intents transactions refunds promotions wallet commissions payouts and subscriptions"),
        ("005", "POS operations", "Printers registers shifts receipts optional service feedback notifications and onboarding"),
        ("006", "Reliability", "Idempotency webhooks outbox audit privacy and login challenges"),
        ("007", "Seeds and constraints", "Default features roles permissions statuses validation functions triggers and indexes"),
    ]
    add_data_table(doc, ["Order", "Layer", "Contents"], rows, [0.65, 1.25, 4.65], font_size=8.8)

    doc.add_heading("Local operating workflow", level=2)
    rows = [
        ("Up", "Starts PostgreSQL 16 Alpine using Docker Compose."),
        ("Migrate", "Applies migrations in order and records their SHA 256 checksum in a ledger."),
        ("Verify", "Runs the transactional schema smoke test and rolls fixture data back."),
        ("Status", "Shows the database and migration state."),
        ("Logs", "Shows container logs for diagnosis."),
        ("Down", "Stops the local database; volume removal is an explicit separate choice."),
    ]
    add_data_table(doc, ["Command", "Purpose"], rows, [1.2, 5.35])

    doc.add_heading("Verification evidence", level=2)
    for item in [
        "All seven migrations replayed successfully from an empty PostgreSQL 16 Docker volume.",
        "A second migration run skipped all files after matching stored SHA 256 checksums.",
        "The smoke test created campus canteen outlet counter terminal catalog variant add-on tax order token transition history and outbox records transactionally.",
        "The smoke test verified idempotent token allocation and rejected an unauthorized status transition.",
        "Every table with updated_at has an automatic timestamp trigger.",
        "High-concurrency load testing and production restore testing remain future release gates.",
    ]:
        doc.add_paragraph(item, style="List Bullet")


def add_readiness(doc: Document) -> None:
    doc.add_heading("Production readiness and TBD decisions", level=1)
    doc.add_paragraph(
        "The schema is production-oriented and locally verified, but production readiness belongs to the complete system. The following work must be completed or approved before launch."
    )
    rows = [
        ("Backend framework and runtime identity", "Select framework session design database roles and whether row-level security is used.", "TBD"),
        ("Student authentication", "Approve phone OTP Truecaller guest access or a subset with abuse controls.", "Blocking for auth flow"),
        ("Payment methods", "Approve launch methods providers split tender behavior settlement and KOT timing.", "Blocking for checkout"),
        ("Tax and rounding", "Approve calculation order inclusive extraction discount interaction and accounting fixtures.", "Blocking for totals"),
        ("Pickup lifecycle", "Confirm seeded transitions cancellation rejection refund and no-show behavior.", "Blocking for UI API"),
        ("Optional features", "Approve wallet coupons inventory feedback dine-in reservation delivery commissions and subscriptions individually.", "Disabled by default"),
        ("Realtime and jobs", "Select event transport queue worker printer adapter retries and monitoring.", "TBD"),
        ("Deployment", "Select hosting database storage CI backups restore monitoring alerts and incident ownership.", "TBD"),
        ("Compliance", "Confirm Indian privacy tax invoice payment campus minor consent retention and deletion requirements.", "TBD owner legal"),
        ("Performance", "Load test campus rush traffic simultaneous tokens payment retries KOT printing and stale POS clients.", "Required before launch"),
    ]
    add_data_table(doc, ["Area", "Required decision or implementation", "Status"], rows, [1.35, 4.0, 1.2], font_size=8.35)


def add_catalogue(doc: Document, tables) -> None:
    doc.add_heading("Complete table catalogue", level=1)
    doc.add_paragraph(
        "This catalogue lists every application table created by migrations 001 through 006. Migration 007 adds seed data functions triggers constraints and indexes but does not add application tables. Key fields omit universal identity and audit columns to keep the reference readable."
    )
    by_domain = {}
    for table in tables:
        by_domain.setdefault(table["domain"], []).append(table)
    for domain, domain_tables in by_domain.items():
        doc.add_heading(domain, level=2)
        rows = []
        for table in domain_tables:
            key_columns = [c for c in table["columns"] if c not in SKIP_COLUMNS][:6]
            refs = table["references"][:5]
            rows.append((
                table["name"],
                purpose_for(table["name"]),
                ", ".join(key_columns) if key_columns else "Metadata only",
                ", ".join(refs) if refs else "None",
            ))
        add_data_table(doc, ["Table", "Purpose", "Representative fields", "References"], rows, [1.55, 2.55, 1.55, 1.05], font_size=8.1)


def add_sources(doc: Document) -> None:
    doc.add_heading("Source and maintenance notes", level=1)
    doc.add_paragraph(
        "The authoritative implementation is the ordered SQL under apps backend database migrations. The Markdown database architecture API security testing decisions and known issues documents explain policy and scope. If a schema migration changes behavior, these references and this handbook should be regenerated in the same change."
    )
    rows = [
        ("Database implementation", "apps/backend/database/migrations"),
        ("Migration runner", "apps/backend/database/scripts/Invoke-Database.ps1"),
        ("Schema smoke test", "apps/backend/database/tests/schema_smoke_test.sql"),
        ("Database guide", "docs/DATABASE.md"),
        ("Architecture", "docs/ARCHITECTURE.md"),
        ("Security", "docs/SECURITY.md"),
        ("Testing", "docs/TESTING.md"),
        ("Decisions and TBDs", "docs/DECISIONS.md and docs/KNOWN_ISSUES.md"),
    ]
    add_data_table(doc, ["Reference", "Project location"], rows, [2.0, 4.55])
    doc.add_paragraph(
        "Document status: database foundation implemented and locally verified on 14 September 2026. Application APIs provider integrations and launch policies are not implied by schema presence and remain governed by the project decision log."
    )


def build_document(project_root: Path, logo_source: Path, output: Path) -> None:
    work_dir = project_root / ".artifacts" / "database-handbook"
    work_dir.mkdir(parents=True, exist_ok=True)
    logo = work_dir / "mealy-logo-cropped.png"
    hierarchy = work_dir / "tenant-hierarchy.png"
    order_flow = work_dir / "order-flow.png"
    split_flow = work_dir / "split-payment.png"
    crop_logo(logo_source, logo)
    create_hierarchy_diagram(hierarchy)
    create_order_diagram(order_flow)
    create_split_diagram(split_flow)

    tables = parse_tables(project_root / "apps" / "backend" / "database" / "migrations")
    if len(tables) != 112:
        raise RuntimeError(f"Expected 112 application tables but parsed {len(tables)}")

    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    core = doc.core_properties
    core.title = "Mealy Database Handbook"
    core.subject = "Mealy PostgreSQL database architecture and implementation reference"
    core.author = "Mealy"
    core.keywords = "Mealy, PostgreSQL, database, POS, student ordering, KOT, tokens, payments"

    add_cover(doc, logo)
    cover_section = doc.sections[0]
    cover_section.header.is_linked_to_previous = False
    cover_section.footer.is_linked_to_previous = False
    cover_section.header.paragraphs[0].text = ""
    cover_section.footer.paragraphs[0].text = ""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    add_header_footer(section, logo)

    add_intro(doc)
    add_metrics(doc)
    add_split_payment(doc, split_flow)
    add_tenant_section(doc, hierarchy)
    add_catalog_section(doc)
    add_order_section(doc, order_flow)
    add_payment_section(doc)
    add_pos_reliability_section(doc)
    add_invariants(doc)
    add_migrations_testing(doc)
    add_readiness(doc)
    add_catalogue(doc, tables)
    add_sources(doc)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the Mealy database handbook DOCX")
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--logo", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build_document(args.project_root.resolve(), args.logo.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
